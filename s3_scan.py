import boto3
import re
import io
import json
import xml.etree.ElementTree as ET
import yaml
import logging
from configparser import ConfigParser
from botocore.exceptions import NoCredentialsError, PartialCredentialsError
from docx import Document
import openpyxl
import csv
from prettytable import PrettyTable
import textwrap
from colorama import init, Fore, Style
import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed

init(autoreset=True)

# Initialize logging
logging.basicConfig(filename='s3_scanner.log', level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

def scan_sensitive_info(data):
    patterns = [
        r'password\s*=\s*["\']?([^"\']+)',
        r'pass\s*=\s*["\']?([^"\']+)',
        r'pw\s*=\s*["\']?([^"\']+)',
        r'cred\s*=\s*["\']?([^"\']+)',
        r'credential\s*=\s*["\']?([^"\']+)',
        r'User Id=([^;]+);Password=([^;]+);',
        r'(A3T[A-Z0-9]|AKIA|AGPA|AROA|AIPA|ANPA|ANVA|ASIA)[A-Z0-9]{16}',
        r'(?<![A-Za-z0-9+/])[A-Za-z0-9+/]{40}(?![A-Za-z0-9+/])',
        r'.{5}~.{34}',
        r'\bpassword\b[\s:]*[\w]+',
        r'\bpass\b[\s:]*[\w]+',
        r'\bpwd\b[\s:]*[\w]+'
    ]

    matches = []
    for pattern in patterns:
        for match in re.finditer(pattern, data, re.IGNORECASE):
            start = max(0, match.start() - 50)
            end = min(len(data), match.end() + 50)
            context = data[start:end]
            matches.append(context)

    return matches

def read_text_file(content):
    return content.decode('utf-8')

def read_word_file(content):
    doc = Document(io.BytesIO(content))
    return '\n'.join([para.text for para in doc.paragraphs])

def read_excel_file(content):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheet = wb.active
    return '\n'.join(['\t'.join([str(cell.value) for cell in row]) for row in sheet.iter_rows()])

def read_csv_file(content):
    content_str = content.decode('utf-8')
    reader = csv.reader(io.StringIO(content_str))
    return '\n'.join([','.join(row) for row in reader])

def read_json_file(content):
    data = json.loads(content.decode('utf-8'))
    return json.dumps(data, indent=4)

def read_xml_file(content):
    tree = ET.ElementTree(ET.fromstring(content.decode('utf-8')))
    root = tree.getroot()
    return ET.tostring(root, encoding='utf8', method='xml').decode('utf-8')

def read_yaml_file(content):
    data = yaml.safe_load(content)
    return yaml.dump(data, default_flow_style=False)

def read_ini_file(content):
    config = ConfigParser()
    config.read_string(content.decode('utf-8'))
    output = io.StringIO()
    config.write(output)
    return output.getvalue()

def read_sql_file(content):
    return content.decode('utf-8')

def process_object(bucket_name, key, content, results):
    try:
        if key.endswith('.txt'):
            data = read_text_file(content)
        elif key.endswith('.docx'):
            data = read_word_file(content)
        elif key.endswith('.xlsx'):
            data = read_excel_file(content)
        elif key.endswith('.csv'):
            data = read_csv_file(content)
        elif key.endswith('.json'):
            data = read_json_file(content)
        elif key.endswith('.xml'):
            data = read_xml_file(content)
        elif key.endswith('.yaml') or key.endswith('.yml'):
            data = read_yaml_file(content)
        elif key.endswith('.ini'):
            data = read_ini_file(content)
        elif key.endswith('.sql'):
            data = read_sql_file(content)
        else:
            logging.info(f"Skipping unsupported file type: {key}")
            print(f"Skipping unsupported file type: {key}")
            return

        sensitive_info = scan_sensitive_info(data)
        if sensitive_info:
            logging.info(f"Sensitive information found in {bucket_name}/{key}")
            print(Fore.YELLOW + f"Sensitive information found in {bucket_name}/{key}:")
            for info in sensitive_info:
                print(info)
                results.append((f"{bucket_name}/{key}", info))
        else:
            logging.info(f"No sensitive information found in {bucket_name}/{key}")
            print(Fore.GREEN + f"No sensitive information found in {bucket_name}/{key}.")
    except Exception as e:
        logging.error(f"Could not process object {key} in bucket {bucket_name}: {e}")
        print(Fore.RED + f"Could not process object {key} in bucket {bucket_name}: {e}")

def read_s3_buckets(bucket_names=None):
    results = []
    try:
        s3 = boto3.client('s3')
        buckets = s3.list_buckets()

        existing_buckets = {bucket['Name'] for bucket in buckets['Buckets']}
        if bucket_names:
            for bucket_name in bucket_names:
                if bucket_name not in existing_buckets:
                    logging.error(f"Bucket {bucket_name} does not exist")
                    print(Fore.RED + f"Bucket {bucket_name} does not exist")
            bucket_names = [bucket for bucket in bucket_names if bucket in existing_buckets]
        else:
            bucket_names = existing_buckets

        with ThreadPoolExecutor() as executor:
            futures = []
            for bucket_name in bucket_names:
                logging.info(f"Scanning bucket: {bucket_name}")
                print(Fore.CYAN + f"Scanning bucket: {bucket_name}")
                
                try:
                    objects = s3.list_objects_v2(Bucket=bucket_name)

                    if 'Contents' not in objects:
                        logging.warning(f"Bucket {bucket_name} is empty or not accessible")
                        print(Fore.YELLOW + f"Bucket {bucket_name} is empty or not accessible.")
                        continue

                    for obj in objects['Contents']:
                        key = obj['Key']
                        logging.info(f"Reading object: {key}")
                        print(f"Reading object: {key}")

                        futures.append(executor.submit(process_object_async, s3, bucket_name, key, results))
                
                except Exception as e:
                    logging.error(f"Could not list objects in bucket {bucket_name}: {e}")
                    print(Fore.RED + f"Could not list objects in bucket {bucket_name}: {e}")

            for future in as_completed(futures):
                future.result()
    
    except (NoCredentialsError, PartialCredentialsError):
        logging.error("Error: AWS credentials not configured properly")
        print(Fore.RED + "Error: AWS credentials not configured properly.")

    return results

def process_object_async(s3, bucket_name, key, results):
    try:
        file_obj = s3.get_object(Bucket=bucket_name, Key=key)
        file_content = file_obj['Body'].read()
        process_object(bucket_name, key, file_content, results)
    except Exception as e:
        logging.error(f"Could not read object {key} in bucket {bucket_name}: {e}")
        print(Fore.RED + f"Could not read object {key} in bucket {bucket_name}: {e}")

def print_results_table(results):
    table = PrettyTable()
    table.field_names = ["File Name", "Sensitive Content"]
    table.align["Sensitive Content"] = "l"

    for file_name, content in results:
        wrapped_content = textwrap.fill(content, width=60)
        table.add_row([file_name, wrapped_content])

    print(Fore.MAGENTA + "\nSensitive Information Summary:")
    print(Fore.CYAN + str(table))

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="S3 Sensitive Information Scanner. ")
    parser.add_argument('-b', '--buckets', nargs='+', help='List of S3 buckets to scan else it will scan through all readable buckets', default=None)
    args = parser.parse_args()

    results = read_s3_buckets(args.buckets)
    print_results_table(results)
